"""
Shared text and image extraction utilities.
Used by gmail/parser.py, telegram_bot/bot.py, agent_sdk/filer.py, and core/agent.py.
"""
import base64
import io
import logging
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from time import monotonic

from core.llmops import record_llm_call
from core.llm_client import create_llm_client, get_model_name
from core.opslog import record_activity, record_issue
from core.structured_output import response_text

logger = logging.getLogger(__name__)


def extract_text(data: bytes, mime_type: str, filename: str) -> str:
    """Extract text content from file bytes based on MIME type."""
    started = monotonic()
    backend = "none"
    text = ""
    try:
        if mime_type == "application/pdf" or filename.lower().endswith(".pdf"):
            text, backend = _extract_pdf_text_with_backend(data)
        elif (
            mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or filename.lower().endswith(".docx")
        ):
            text = extract_docx_text(data)
            backend = "docx"
        elif mime_type.startswith("text/"):
            text = data.decode("utf-8", errors="replace")
            backend = "text"
        elif (
            mime_type in (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel",
            )
            or filename.lower().endswith((".xlsx", ".xls"))
        ):
            text = extract_spreadsheet_text(data, filename)
            backend = "spreadsheet"
        elif filename.lower().endswith(".csv"):
            text = extract_spreadsheet_text(data, filename)
            backend = "csv"
    except Exception as e:
        logger.warning("Text extraction failed for %s: %s", filename, e)
        record_issue(
            level="WARNING",
            event="document_text_extraction_failed",
            component="document",
            status="warning",
            summary="Document text extraction failed",
            metadata={"filename": filename[:120], "mime_type": mime_type, "error": str(e)},
        )
    record_activity(
        event="document_text_extracted",
        component="document",
        summary="Extracted local document text for downstream processing",
        duration_ms=(monotonic() - started) * 1000,
        metadata={
            "filename": filename[:120],
            "mime_type": mime_type,
            "backend": backend,
            "extracted_chars": len(text),
            "text_nonempty": bool(text.strip()),
        },
    )
    return text


def extract_pdf_text(data: bytes) -> str:
    text, _backend = _extract_pdf_text_with_backend(data)
    return text


def _extract_pdf_text_with_backend(data: bytes) -> tuple[str, str]:
    primary_text = _extract_pdf_text_pypdf(data)
    if len(primary_text.strip()) >= 80:
        return primary_text, "pdf_pypdf2"

    ocr_text = _ocr_pdf_first_page(data)
    if len(ocr_text.strip()) > len(primary_text.strip()):
        return ocr_text, "pdf_ocr_tesseract"
    return primary_text, "pdf_pypdf2"


def _extract_pdf_text_pypdf(data: bytes) -> str:
    import PyPDF2
    reader = PyPDF2.PdfReader(io.BytesIO(data), strict=False)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n".join(pages)


def _ocr_pdf_first_page(data: bytes) -> str:
    sips_bin = shutil.which("sips")
    tesseract_bin = shutil.which("tesseract")
    if not sips_bin or not tesseract_bin:
        return ""

    with tempfile.TemporaryDirectory() as temp_dir:
        pdf_path = Path(temp_dir) / "input.pdf"
        png_path = Path(temp_dir) / "input.png"
        txt_base = Path(temp_dir) / "ocr"
        pdf_path.write_bytes(data)
        try:
            subprocess.run(
                [sips_bin, "-s", "format", "png", str(pdf_path), "--out", str(png_path)],
                check=True,
                capture_output=True,
                text=True,
            )
            subprocess.run(
                [tesseract_bin, str(png_path), str(txt_base), "--psm", "6"],
                check=True,
                capture_output=True,
                text=True,
            )
        except Exception as exc:
            logger.warning("PDF OCR fallback failed: %s", exc)
            return ""

        txt_path = Path(f"{txt_base}.txt")
        if not txt_path.exists():
            return ""
        return txt_path.read_text(encoding="utf-8", errors="replace").strip()


def extract_docx_text(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_spreadsheet_text(data: bytes, filename: str) -> str:
    if filename.lower().endswith(".csv"):
        return data.decode("utf-8", errors="replace")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    lines.append(row_text)
        return "\n".join(lines)
    except Exception as e:
        logger.warning("Spreadsheet extraction failed for %s: %s", filename, e)
        return ""


def describe_image(image_data: bytes, mime_type: str) -> str:
    """
    Use the configured multimodal model to describe an image.
    Returns a text description including any visible text, amounts, dates, names.
    """
    try:
        b64 = base64.standard_b64encode(image_data).decode("utf-8")
        client = create_llm_client()
        model_name = get_model_name("vision")
        started_at = datetime.now(timezone.utc).isoformat()
        started_clock = monotonic()
        response = client.messages.create(
            model=model_name,
            max_tokens=1024,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "image",
                            "source": {
                                "type": "base64",
                                "media_type": mime_type,
                                "data": b64,
                            },
                        },
                        {
                            "type": "text",
                            "text": (
                                "Describe this image. If it is a document, extract all visible text. "
                                "Note any key details such as amounts, dates, names, addresses, "
                                "or document type. Be thorough."
                            ),
                        },
                    ],
                }
            ],
        )
        record_llm_call(
            task="vision",
            model=model_name,
            status="ok",
            started_at=started_at,
            latency_ms=(monotonic() - started_clock) * 1000,
            response=response,
            metadata={"channel": "vision", "mime_type": mime_type},
        )
        return response_text(response)[:4000]
    except Exception as e:
        model_name = locals().get("model_name") or get_model_name("vision")
        record_llm_call(
            task="vision",
            model=model_name,
            status="api_error",
            started_at=locals().get("started_at"),
            latency_ms=(monotonic() - locals().get("started_clock", monotonic())) * 1000,
            error=str(e),
            metadata={"channel": "vision", "mime_type": mime_type},
        )
        logger.warning("Image description failed: %s", e)
        return ""
